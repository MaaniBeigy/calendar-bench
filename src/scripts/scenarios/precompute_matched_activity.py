"""Resolve curated HealthTask matched_activity labels to HumanActivities IRIs.

Reads the curator-authored xlsx and the HumanActivities TTL, then emits
one JSONL record per HealthTask with the resolved activity IRI(s).
Strict resolution: any unresolved label aborts the run.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
from dataclasses import dataclass
from difflib import get_close_matches
from pathlib import Path
from typing import Iterable

from rdflib import Graph, Namespace
from rdflib.namespace import RDFS

HA_INSTANCE_PREFIX = "https://w3id.org/calendar-bench/human-activities/activity/"
HB_TASK_PREFIX = "https://w3id.org/calendar-bench/health/task/"

_HA = Namespace("https://w3id.org/calendar-bench/human-activities/")
_HA_ACT = Namespace(HA_INSTANCE_PREFIX)
_HB_TK = Namespace(HB_TASK_PREFIX)


@dataclass(frozen=True, slots=True)
class ResolvedRow:
    """One xlsx row resolved to a HealthTask IRI and matched activity IRIs."""

    task_slug: str
    task_iri: str
    matched_activity_labels: tuple[str, ...]
    matched_activity_iris: tuple[str, ...]
    note: str


@dataclass(frozen=True, slots=True)
class UnresolvedRow:
    """One xlsx row whose label did not match any HumanActivities entry."""

    task_slug: str
    raw_label: str
    nearest_matches: tuple[str, ...]


class LabelResolutionError(RuntimeError):
    """Raised when one or more xlsx labels do not match the loaded ontology."""


def _normalize(text: str) -> str:
    """Lowercase, collapse whitespace, normalize dashes for label compare."""
    text = text.replace("–", "-").replace("—", "-")
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def load_label_index(ttl_path: Path) -> dict[str, str]:
    """Return `{normalized rdfs:label: IRI}` for every ha-act:* instance.

    Only nodes under the HumanActivities activity prefix are indexed.
    """
    graph = Graph()
    graph.parse(ttl_path.as_posix(), format="turtle")
    index: dict[str, str] = {}
    for subject, _, label in graph.triples((None, RDFS.label, None)):
        iri = str(subject)
        if not iri.startswith(HA_INSTANCE_PREFIX):
            continue
        index[_normalize(str(label))] = iri
    return index


def _slugify(task_name: str) -> str:
    """Lower-kebab-case slug used as the HealthTask local-name."""
    text = task_name.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


def _strip_emoji_suffix(task_name: str) -> str:
    """Drop trailing non-ASCII glyphs commonly used as decoration."""
    return re.sub(r"\s+[^\w\s,.:;()'\"-]+\s*$", "", task_name).strip()


def _split_multi(raw_label: str) -> list[str]:
    """Split a matched_activity cell on the multi-target separator.

    Multi-target separator is the literal `||` so existing single labels
    that contain semicolons or commas inside their HumanActivities text
    are not falsely split. v1 curation has no multi-target rows.
    """
    parts = [p.strip() for p in raw_label.split("||")]
    return [p for p in parts if p]


def _read_xlsx_rows(xlsx_path: Path) -> list[dict[str, str]]:
    """Return the xlsx rows as a list of dicts keyed by column header."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path.as_posix(), read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c) if c is not None else "" for c in rows[0]]
    out: list[dict[str, str]] = []
    for raw in rows[1:]:
        record = {
            header[i]: ("" if raw[i] is None else str(raw[i]))
            for i in range(len(header))
        }
        out.append(record)
    return out


def resolve_rows(
    xlsx_path: Path,
    ttl_path: Path,
) -> tuple[list[ResolvedRow], list[UnresolvedRow]]:
    """Resolve every xlsx row's matched_activity label to a HumanActivities IRI."""
    label_index = load_label_index(ttl_path)
    rows = _read_xlsx_rows(xlsx_path)
    resolved: list[ResolvedRow] = []
    unresolved: list[UnresolvedRow] = []
    for row in rows:
        task_name = _strip_emoji_suffix(row.get("task_name", ""))
        if not task_name:
            continue
        raw_label = row.get("matched_activity", "").strip()
        note = row.get("matched_activity_note", "").strip()
        slug = _slugify(task_name)
        task_iri = f"{HB_TASK_PREFIX}{slug}"
        if not raw_label:
            unresolved.append(
                UnresolvedRow(task_slug=slug, raw_label="", nearest_matches=())
            )
            continue
        parts = _split_multi(raw_label)
        labels: list[str] = []
        iris: list[str] = []
        row_failed = False
        for part in parts:
            iri = label_index.get(_normalize(part))
            if iri is None:
                nearest = get_close_matches(_normalize(part), label_index.keys(), n=3)
                unresolved.append(
                    UnresolvedRow(
                        task_slug=slug,
                        raw_label=part,
                        nearest_matches=tuple(nearest),
                    )
                )
                row_failed = True
                continue
            labels.append(part)
            iris.append(iri)
        if row_failed or not iris:
            continue
        resolved.append(
            ResolvedRow(
                task_slug=slug,
                task_iri=task_iri,
                matched_activity_labels=tuple(labels),
                matched_activity_iris=tuple(iris),
                note=note,
            )
        )
    return resolved, unresolved


def _format_unresolved(unresolved: Iterable[UnresolvedRow]) -> str:
    lines = ["Unresolved matched_activity labels:"]
    for row in unresolved:
        if not row.raw_label:
            lines.append(f"  task={row.task_slug}: (empty cell)")
            continue
        suggest = (
            f"; nearest: {', '.join(row.nearest_matches)}"
            if row.nearest_matches
            else ""
        )
        lines.append(f"  task={row.task_slug}: {row.raw_label!r}{suggest}")
    return "\n".join(lines)


def write_jsonl(rows: Iterable[ResolvedRow], out_path: Path) -> int:
    """Write one JSON record per row; return record count."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    timestamp = _dt.datetime.now(_dt.timezone.utc).isoformat(timespec="seconds")
    count = 0
    with out_path.open("w", encoding="utf-8") as fh:
        for row in rows:
            record = {
                "task_slug": row.task_slug,
                "task_iri": row.task_iri,
                "matched_activity_labels": list(row.matched_activity_labels),
                "matched_activity_iris": list(row.matched_activity_iris),
                "note": row.note,
                "ts": timestamp,
            }
            fh.write(json.dumps(record, ensure_ascii=False) + "\n")
            count += 1
    return count


def run(
    xlsx_path: Path,
    ttl_path: Path,
    out_path: Path,
) -> int:
    """End-to-end: resolve rows and write the JSONL. Aborts on unresolved labels."""
    resolved, unresolved = resolve_rows(xlsx_path, ttl_path)
    if unresolved:
        raise LabelResolutionError(_format_unresolved(unresolved))
    return write_jsonl(resolved, out_path)


def _parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--ha-ttl", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv)
    try:
        count = run(args.xlsx, args.ha_ttl, args.out)
    except LabelResolutionError as exc:
        sys.stderr.write(str(exc) + "\n")
        return 2
    sys.stdout.write(f"resolved {count} tasks to {args.out}\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
